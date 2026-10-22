"""Import an existing DOCX/PDF resume into the structured master resume."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader

from job_applier.llm import LLM
from job_applier.pipeline import ai
from job_applier.profile.models import Profile, Resume


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(line for line in lines if line.strip())
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8")
    raise ValueError(f"Unsupported resume format {suffix!r}; use .docx, .pdf or .txt")


def import_resume(llm: LLM, path: Path) -> Resume:
    text = extract_text(path)
    if len(text.strip()) < 200:
        raise ValueError("Very little text found; is the PDF a scanned image? Try the DOCX instead.")
    return ai.parse_resume(llm, text)


def prefill_profile(profile: Profile, resume: Resume) -> Profile:
    """Copy obvious facts from the resume into empty profile fields."""
    out = profile.model_copy(deep=True)
    parts = resume.basics.name.split()
    if parts and not out.first_name:
        out.first_name = parts[0]
    if len(parts) > 1 and not out.last_name:
        out.last_name = parts[-1]
    out.email = out.email or resume.basics.email
    out.phone = out.phone or resume.basics.phone
    if resume.basics.location and not out.city:
        loc = [p.strip() for p in resume.basics.location.split(",")]
        out.city = loc[0]
        if len(loc) > 1 and not out.state:
            out.state = loc[1]
    for link in resume.basics.links:
        url = link.url.lower()
        if "linkedin" in url and not out.linkedin_url:
            out.linkedin_url = link.url
        elif "github" in url and not out.github_url:
            out.github_url = link.url
        elif not out.website_url and "linkedin" not in url and "github" not in url:
            out.website_url = link.url
    if resume.experience:
        out.current_company = out.current_company or resume.experience[0].company
        out.current_title = out.current_title or resume.experience[0].title
    return out
