"""Render the structured resume (and cover letters) to PDF and DOCX."""

from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from jinja2 import Environment, PackageLoader, select_autoescape

from job_applier.browser import launch_browser
from job_applier.profile.models import Resume

_env = Environment(loader=PackageLoader("job_applier.resume", "templates"), autoescape=select_autoescape())


def resume_html(resume: Resume) -> str:
    return _env.get_template("resume.html").render(r=resume)


def file_stem(name: str, suffix: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_") or "Candidate"
    return f"{safe}_{suffix}"


def html_to_pdf(content: str, out: Path) -> Path:
    out.parent.mkdir(parents=True, exist_ok=True)
    with launch_browser(headless=True) as browser:
        page = browser.new_page()
        page.set_content(content, wait_until="load")
        page.pdf(path=str(out), format="Letter", print_background=True, prefer_css_page_size=True)
    return out


def render_pdf(resume: Resume, out: Path) -> Path:
    return html_to_pdf(resume_html(resume), out)


def cover_letter_pdf(text: str, out: Path) -> Path:
    paragraphs = "".join(f"<p>{html.escape(p)}</p>" for p in text.split("\n\n") if p.strip())
    page = (
        "<!doctype html><html><head><meta charset='utf-8'><style>@page{size:Letter;margin:1in}"
        "body{font-family:Calibri,'Segoe UI',Arial,sans-serif;font-size:11pt;line-height:1.45}"
        f"p{{margin:0 0 10pt;white-space:pre-line}}</style></head><body>{paragraphs}</body></html>"
    )
    return html_to_pdf(page, out)


def render_docx(resume: Resume, out: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    doc.add_heading(resume.basics.name, level=0)
    contact = [resume.basics.location, resume.basics.phone, resume.basics.email]
    contact += [link.url for link in resume.basics.links]
    doc.add_paragraph(" | ".join(c for c in contact if c))
    if resume.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(resume.summary)
    if resume.experience:
        doc.add_heading("Experience", level=1)
        for exp in resume.experience:
            p = doc.add_paragraph()
            p.add_run(f"{exp.title}, {exp.company}").bold = True
            p.add_run(f"    {exp.start} – {exp.end or 'Present'}")
            for bullet in exp.bullets:
                doc.add_paragraph(bullet, style="List Bullet")
    if resume.skills:
        doc.add_heading("Skills", level=1)
        for group in resume.skills:
            p = doc.add_paragraph()
            p.add_run(f"{group.name}: ").bold = True
            p.add_run(", ".join(group.items))
    if resume.projects:
        doc.add_heading("Projects", level=1)
        for proj in resume.projects:
            p = doc.add_paragraph()
            p.add_run(proj.name).bold = True
            if proj.description:
                doc.add_paragraph(proj.description)
            for bullet in proj.bullets:
                doc.add_paragraph(bullet, style="List Bullet")
    if resume.education:
        doc.add_heading("Education", level=1)
        for edu in resume.education:
            p = doc.add_paragraph()
            p.add_run(", ".join(x for x in (edu.degree, edu.field) if x)).bold = True
            p.add_run(f"    {edu.institution} {edu.end}".rstrip())
    if resume.certifications:
        doc.add_heading("Certifications", level=1)
        for cert in resume.certifications:
            doc.add_paragraph(cert, style="List Bullet")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
